"""선량 평가 부품 목록. components.json + (있으면) 소재 파트가 채운 component_inputs.xlsx."""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

from . import CONFIG, config


@dataclass
class Component:
    id: str
    name: str
    group: str
    link: int
    offset: np.ndarray
    detail: str = ""
    reference: bool = False
    S: float = 1.0                 # 차폐 투과율 (0~1)
    tid_max: float = math.nan      # 허용 누적선량 [Gy]
    shield_mass: float = 0.0       # 차폐재 질량 [kg]
    source_note: str = ""
    extra: dict = field(default_factory=dict)


def _num(v):
    try:
        if v is None or (isinstance(v, str) and not v.strip()):
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def load(json_path=None, xlsx_path=None) -> tuple[list[Component], dict]:
    """부품 목록과 설정값(N_ex 등)을 읽는다. xlsx가 있으면 S, TID_max, 질량을 덮어쓴다."""
    cfg = config.components_config(json_path)
    comps = [Component(id=c["id"], name=c["name"], group=c["group"], link=int(c["link"]),
                       offset=np.asarray(c["offset"], float), detail=c.get("detail", ""),
                       reference=bool(c.get("reference", False))) for c in cfg["components"]]
    settings = {"N_ex": cfg.get("N_ex_default")}
    xlsx_path = Path(xlsx_path) if xlsx_path else CONFIG / "component_inputs.xlsx"
    if xlsx_path.exists():
        _apply_xlsx(comps, settings, xlsx_path)
    return comps, settings


def _apply_xlsx(comps: list[Component], settings: dict, path: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb["부품_입력"]
    header_row = None
    for r in range(1, 10):
        if ws.cell(r, 1).value == "ID":
            header_row = r
            break
    if header_row is None:
        raise ValueError("부품_입력 시트에서 'ID' 머리글을 찾지 못했습니다")
    cols = {str(ws.cell(header_row, c).value).split("\n")[0].strip(): c
            for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value}
    by_id = {c.id: c for c in comps}

    def get(r, key):
        c = cols.get(key)
        return ws.cell(r, c).value if c else None

    for r in range(header_row + 1, ws.max_row + 1):
        cid = get(r, "ID")
        if cid not in by_id:
            continue
        comp = by_id[cid]
        s_direct = _num(get(r, "S_k 직접입력"))
        mu, rho, x = _num(get(r, "μ/ρ")), _num(get(r, "ρ")), _num(get(r, "두께 x"))
        if s_direct is not None:
            comp.S = s_direct
        elif None not in (mu, rho, x):
            comp.S = math.exp(-mu * rho * x)
        tid = _num(get(r, "TID_max"))
        if tid is not None:
            comp.tid_max = tid
        m = _num(get(r, "차폐재 질량"))
        if m is not None:
            comp.shield_mass = m
        comp.source_note = str(get(r, "근거·출처") or "")
        if not 0.0 <= comp.S <= 1.0:
            raise ValueError(f"{cid}: 투과율 S_k={comp.S} 가 0~1 범위를 벗어났습니다")

    if "설정" in wb.sheetnames:
        ws2 = wb["설정"]
        for r in range(1, ws2.max_row + 1):
            if ws2.cell(r, 1).value == "N_ex":
                settings["N_ex"] = _num(ws2.cell(r, 3).value)


def budgets(comps: list[Component], N_ex) -> np.ndarray:
    """1회 작업당 선량 예산 B_k = TID_max,k / N_ex (값이 없으면 nan)."""
    if not N_ex:
        return np.full(len(comps), np.nan)
    return np.array([c.tid_max / N_ex if not math.isnan(c.tid_max) else np.nan for c in comps])
