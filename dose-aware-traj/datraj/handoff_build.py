"""인계 파일 생성 도우미: 부품 위치도, 소재 입력 xlsx, 제어용 2관절 축약 모델."""
from __future__ import annotations

from pathlib import Path

import numpy as np

from . import viz_style
from .components import Component


# ---------------------------------------------------------------------------
# 부품 위치도
# ---------------------------------------------------------------------------

def component_positions(robot, comps: list[Component], q) -> np.ndarray:
    return np.array([robot.point_on_link(q, c.link, c.offset) for c in comps])


def draw_component_figure(robot, comps: list[Component], q, out: Path) -> None:
    viz_style.apply()
    import matplotlib.pyplot as plt
    from matplotlib import patheffects as pe

    pos, R = robot.link_frames(q)
    flange = pos[6] + R[6] @ np.array([0, 0, robot.flange])
    tcp = pos[6] + R[6] @ robot.tcp_local
    chain = np.vstack([robot.base_pos, pos, flange])
    pts = component_positions(robot, comps, q)

    fig, ax = plt.subplots(figsize=(7.2, 6.4))
    ax.plot(chain[:, 0], chain[:, 2], color=viz_style.INK_2, lw=6, alpha=0.25, solid_capstyle="round")
    ax.plot(chain[:, 0], chain[:, 2], color=viz_style.INK_2, lw=1.2)
    ax.plot([flange[0], tcp[0]], [flange[2], tcp[2]], color=viz_style.MUTED, lw=4, alpha=0.5,
            solid_capstyle="round")
    ax.plot(pos[:, 0], pos[:, 2], "o", ms=5, mfc=viz_style.SURFACE, mec=viz_style.INK_2, mew=1.2)
    ax.add_patch(plt.Rectangle((-0.11, -0.02), 0.22, 0.02, color=viz_style.AXIS))
    halo = [pe.withStroke(linewidth=3, foreground=viz_style.SURFACE)]
    groups = {"drive": viz_style.CATEGORICAL[0], "tool": viz_style.CATEGORICAL[1],
              "sensor": viz_style.CATEGORICAL[2], "base": viz_style.CATEGORICAL[3],
              "reference": viz_style.MUTED}
    offsets = {"C01": (-0.13, 0.0), "C02": (-0.13, 0.0), "C03": (-0.13, 0.02), "C04": (0.05, 0.03),
               "C05": (0.05, 0.03), "C06": (0.05, 0.03), "C07": (0.06, 0.04), "C08": (0.07, 0.0),
               "C09": (-0.05, -0.04), "C10": (-0.05, -0.03), "P0": (0.05, -0.05)}
    for c, pt in zip(comps, pts):
        col = groups.get(c.group, viz_style.INK)
        ax.plot(pt[0], pt[2], marker="o" if not c.reference else "D", ms=8, color=col, mec=viz_style.SURFACE,
                mew=1.5, zorder=5)
        dx, dz = offsets.get(c.id, (0.05, 0.0))
        ax.annotate(f"{c.id} {c.name}", (pt[0], pt[2]), (pt[0] + dx, pt[2] + dz), fontsize=8,
                    color=viz_style.INK, va="center", ha="left" if dx > 0 else "right",
                    path_effects=halo, zorder=6)
    ax.set_aspect("equal")
    ax.set_xlabel("x [m]")
    ax.set_ylabel("z [m]")
    ax.set_title("선량 평가점 위치 (옆에서 본 모습, 예시 자세)", loc="left")
    ax.text(0.0, -0.16, "관절 구동부(C01~C07)는 각 관절 축 중심에 둔다. 위치는 components.json의 link, offset으로 정의.",
            transform=ax.transAxes, fontsize=8, color=viz_style.INK_2)
    xs = np.concatenate([chain[:, 0], pts[:, 0], [tcp[0]]])
    zs = np.concatenate([chain[:, 2], pts[:, 2], [tcp[2]]])
    ax.set_xlim(xs.min() - 0.35, xs.max() + 0.45)
    ax.set_ylim(-0.05, zs.max() + 0.12)
    ax.grid(True)
    fig.savefig(out)
    plt.close(fig)


# ---------------------------------------------------------------------------
# 소재 입력 xlsx
# ---------------------------------------------------------------------------

def build_materials_xlsx(comps: list[Component], out: Path, figure: Path | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    title = Font(name=F, size=13, bold=True)
    bold = Font(name=F, size=10, bold=True)
    normal = Font(name=F, size=10)
    note = Font(name=F, size=9, color="52514E")
    example = Font(name=F, size=10, italic=True, color="898781")
    head_fill = PatternFill("solid", fgColor="EDEDEA")
    input_fill = PatternFill("solid", fgColor="FFF2B3")
    fixed_fill = PatternFill("solid", fgColor="F5F5F3")
    thin = Side(style="thin", color="C3C2B7")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()

    # ---- 안내
    ws = wb.active
    ws.title = "안내"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 90
    ws["B1"] = "부품별 차폐·내방사선 입력 양식"
    ws["B1"].font = title
    ws["B2"] = "소재 파트 → 경로계획(E) 파트  |  양식 v0.1  |  2026-10-01"
    ws["B2"].font = note
    rows = [
        ("목적", "로봇 부품 k마다 누적선량 D_k = S_k × ∫Ḋ(p_k(t)) dt 를 계산하고, 1회 작업 예산 B_k = TID_max,k / N_ex 와 비교한다. "
                 "이 표의 S_k, TID_max, 차폐재 질량이 계산에 그대로 들어간다."),
        ("작성할 칸", "노란 칸만 채운다. 회색 칸(ID, 부품, 위치)은 경로계획 쪽 정의이므로 바꾸지 말고, 바꿔야 하면 비고에 적어 달라."),
        ("투과율 S_k", "차폐재·두께·μ/ρ·ρ 를 넣으면 S_k = exp(−(μ/ρ)·ρ·x) 가 자동 계산된다. "
                      "몬테카를로 등으로 따로 구한 값이 있으면 'S_k 직접입력'에 넣는다 (직접입력이 우선). 차폐가 없으면 비워 두면 1로 처리된다."),
        ("μ/ρ 기준 에너지", "설정 시트의 '기준 감마 에너지'와 같은 에너지의 값을 쓴다. 이 에너지는 방사선 파트와 맞춘다."),
        ("TID_max", "부품이 기능을 유지하는 최대 누적 흡수선량 [Gy]. 매질은 설정 시트의 '선량 매질'(기본 Si). 근거(논문, 데이터시트, 시험 결과)를 반드시 적는다."),
        ("차폐재 질량", "해당 부품에 추가되는 차폐재 질량 [kg]. 경로계획 쪽 토크 계산에서 이 부품 위치의 점질량으로 더한다."),
        ("N_ex", "로봇 수명 동안 예상되는 베드 교체 횟수. 공정 파트와 협의해 설정 시트에 넣는다."),
        ("예시 행", "부품_입력 시트 5행은 형식을 보여주기 위한 예시이며 계산에 쓰이지 않는다."),
        ("부품 위치", "부품_위치도 시트 참고. 관절 구동부는 각 관절 축 중심 위치로 둔다."),
        ("회신", "채운 파일을 그대로 보내 주면 된다. 파일 이름은 component_inputs.xlsx 로 유지."),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(i, 2, k).font = bold
        c = ws.cell(i, 3, v)
        c.font = normal
        c.alignment = wrap
        ws.row_dimensions[i].height = 32
    r = 4 + len(rows) + 1
    ws.cell(r, 2, "색 구분").font = bold
    ws.cell(r + 1, 2, "입력 칸").fill = input_fill
    ws.cell(r + 1, 3, "소재 파트가 채우는 칸").font = normal
    ws.cell(r + 2, 2, "고정 칸").fill = fixed_fill
    ws.cell(r + 2, 3, "경로계획 쪽 정의 (수정하지 않음)").font = normal
    ws.cell(r + 3, 2, "계산 칸").font = normal
    ws.cell(r + 3, 3, "수식으로 자동 계산 (흰 바탕)").font = normal
    for rr in range(r + 1, r + 4):
        ws.cell(rr, 2).border = box
        ws.cell(rr, 2).font = normal

    # ---- 부품_입력
    ws = wb.create_sheet("부품_입력")
    headers = [
        ("ID", 7), ("부품", 16), ("설명", 30), ("분류", 9), ("링크\n(pybullet)", 9),
        ("위치\n(링크 좌표계, m)", 18), ("차폐재", 10), ("두께 x\n(cm)", 9), ("μ/ρ\n(cm²/g)", 10),
        ("ρ\n(g/cm³)", 9), ("S_k 계산\n(자동)", 10), ("S_k 직접입력\n(선택)", 12), ("적용 S_k\n(자동)", 10),
        ("차폐재 질량\n(kg)", 11), ("TID_max\n(Gy)", 11), ("1회 예산 B_k\n(Gy, 자동)", 13),
        ("근거·출처", 34), ("비고", 24),
    ]
    ws["A1"] = "부품별 차폐·내방사선 입력"
    ws["A1"].font = title
    ws["A2"] = "노란 칸을 채운다. 5행은 예시(계산 제외). S_k = exp(−(μ/ρ)·ρ·x), B_k = TID_max / N_ex (N_ex는 설정 시트)."
    ws["A2"].font = note
    HR = 4
    for j, (h, w) in enumerate(headers, start=1):
        c = ws.cell(HR, j, h)
        c.font = bold
        c.fill = head_fill
        c.alignment = center
        c.border = box
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[HR].height = 34

    group_ko = {"drive": "구동부", "tool": "공구", "sensor": "센서", "base": "베이스"}

    def put_formulas(r):
        ws.cell(r, 11, f'=IF(COUNT(H{r}:J{r})=3,EXP(-I{r}*J{r}*H{r}),"")')
        ws.cell(r, 13, f'=IF(ISNUMBER(L{r}),L{r},IF(ISNUMBER(K{r}),K{r},1))')
        ws.cell(r, 16, f'=IF(AND(ISNUMBER(O{r}),ISNUMBER(설정!$C$4)),IF(설정!$C$4>0,O{r}/설정!$C$4,""),"")')
        for col in (11, 13, 16):
            ws.cell(r, col).number_format = "0.000"

    # 예시 행
    er = HR + 1
    ex = ["예시", "예: A6 구동부", "형식 예시. 계산에 쓰이지 않음", "구동부", 5, "0, 0, 0", "Pb", 1.0, 0.1614, 11.35,
          None, None, None, 2.0, 1000, None, "예시 값. μ/ρ는 NIST 납 0.5 MeV 값", ""]
    for j, v in enumerate(ex, start=1):
        if v is not None:
            ws.cell(er, j, v)
        ws.cell(er, j).font = example
        ws.cell(er, j).border = box
        ws.cell(er, j).alignment = wrap
    put_formulas(er)
    for col in (11, 13, 16):
        ws.cell(er, col).font = example
    ws.cell(er, 9).comment = Comment(
        "NIST X-Ray Mass Attenuation Coefficients, Lead (Z=82), 0.5 MeV: μ/ρ = 1.614E-01 cm²/g. 예시 행 전용.", "E part")

    real = [c for c in comps if not c.reference]
    for i, comp in enumerate(real):
        r = er + 1 + i
        vals = [comp.id, comp.name, comp.detail, group_ko.get(comp.group, comp.group), comp.link,
                ", ".join(f"{v:g}" for v in comp.offset)]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(r, j, v)
            c.font = normal
            c.fill = fixed_fill
            c.alignment = wrap
            c.border = box
        for col in (7, 8, 9, 10, 12, 14, 15, 17, 18):
            c = ws.cell(r, col)
            c.fill = input_fill
            c.font = normal
            c.border = box
            c.alignment = wrap
        put_formulas(r)
        for col in (11, 13, 16):
            ws.cell(r, col).font = normal
            ws.cell(r, col).border = box
        ws.row_dimensions[r].height = 30
    last = er + len(real)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        showErrorMessage=True, errorTitle="범위 오류", error="투과율은 0~1 사이 값입니다")
    ws.add_data_validation(dv)
    dv.add(f"L{er + 1}:L{last}")
    nn = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="범위 오류", error="0 이상의 숫자를 입력하세요")
    ws.add_data_validation(nn)
    for col in "HIJNO":
        nn.add(f"{col}{er + 1}:{col}{last}")
    ws.freeze_panes = ws.cell(HR + 1, 3)

    # ---- 설정
    ws = wb.create_sheet("설정")
    ws["A1"] = "설정값"
    ws["A1"].font = title
    for j, (h, w) in enumerate([("항목", 22), ("설명", 52), ("값", 12), ("단위", 8), ("담당", 16)], start=1):
        c = ws.cell(3, j, h)
        c.font = bold
        c.fill = head_fill
        c.border = box
        c.alignment = center
        ws.column_dimensions[c.column_letter].width = w
    settings = [
        ("N_ex", "로봇 수명 동안 예상되는 베드 교체 횟수", None, "회", "공정·소재 협의"),
        ("기준 감마 에너지", "μ/ρ를 고를 때 쓰는 광자 에너지. 방사선 파트와 같은 값", None, "MeV", "방사선·소재"),
        ("선량 매질", "TID_max와 선량 격자의 흡수선량 기준 매질", "Si", "-", "소재"),
    ]
    for i, row in enumerate(settings, start=4):
        for j, v in enumerate(row, start=1):
            c = ws.cell(i, j, v)
            c.font = normal
            c.border = box
            c.alignment = wrap
        ws.cell(i, 3).fill = input_fill

    # ---- 부품_위치도
    ws = wb.create_sheet("부품_위치도")
    ws["A1"] = "선량 평가점 위치"
    ws["A1"].font = title
    ws["A2"] = "예시 자세에서 옆에서 본 모습. 실제 계산은 매 시점의 관절각으로 위치를 다시 구한다."
    ws["A2"].font = note
    if figure and Path(figure).exists():
        img = XLImage(str(figure))
        scale = 620 / img.width
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws.add_image(img, "A4")

    wb.active = 1
    wb.save(out)


# ---------------------------------------------------------------------------
# 제어 파트용 2관절 축약 모델
# ---------------------------------------------------------------------------

def reduced_2dof(robot, params: dict) -> dict:
    """A1=A3=A5=A6=A7=0 기준으로 A2(어깨 피치)와 A4(팔꿈치 피치) 2관절 평면 모델의 집중 파라미터."""
    q0 = np.zeros(7)
    pos, R = robot.link_frames(q0)
    links = params["robot"]["links"]
    n = R[1][:, 2]                        # A2 축 방향
    n4 = R[3][:, 2]
    if abs(abs(n @ n4) - 1) > 1e-6:
        raise RuntimeError("기준 자세에서 A2, A4 축이 평행하지 않습니다")
    arm = np.array([0.0, 0.0, 1.0])       # 기준 자세에서 팔이 뻗는 방향
    lateral = np.cross(n, arm)

    def bodies(idx_list, extra=()):
        out = []
        for j in idx_list:
            L = links[j]
            c = pos[j] + R[j] @ np.asarray(L["com"])
            I = R[j] @ np.diag(L["inertia_diag"]) @ R[j].T
            out.append((L["link"], L["mass"], c, I))
        out.extend(extra)
        return out

    def lump(bs, joint_pos):
        m = sum(b[1] for b in bs)
        c = sum(b[1] * b[2] for b in bs) / m
        Ic = 0.0
        for _, mi, ci, Ii in bs:
            d = ci - c
            Ic += n @ Ii @ n + mi * (d @ d - (d @ n) ** 2)
        rel = c - joint_pos
        return {"mass": round(float(m), 4), "com_along_link": round(float(rel @ arm), 5),
                "com_lateral": round(float(rel @ lateral), 5), "com_out_of_plane": round(float(rel @ n), 5),
                "inertia_about_com": round(float(Ic), 6), "bodies": [b[0] for b in bs]}

    flange = pos[6] + R[6] @ np.array([0, 0, robot.flange])
    tcp = pos[6] + R[6] @ robot.tcp_local
    tool_c = pos[6] + R[6] @ robot.tool_mass.local_pos
    tool = ("tool", robot.tool_mass.mass, tool_c, np.zeros((3, 3)))
    seg1 = lump(bodies([1, 2]), pos[1])
    seg2 = lump(bodies([3, 4, 5, 6], extra=[tool]), pos[3])
    bed_off = params["payload"]["bed_com_from_tcp"]
    bed_c = tcp + R[6][:, 2] * bed_off
    return {
        "note": "제어 파트 Simscape 2관절 모델용. 7축 URDF와 공통 파라미터에서 계산한 값 (datraj.handoff_build.reduced_2dof).",
        "reference_configuration": "A1=A3=A5=A6=A7=0 (팔을 위로 곧게 편 자세에서 집중화). 손목을 굽힌 자세 기준값이 필요하면 요청.",
        "joints": {"q1": "A2 어깨 피치", "q2": "A4 팔꿈치 피치"},
        "A4_axis_sign_relative_to_A2": int(round(float(n @ n4))),
        "angle_convention": ("q1=q2=0 이면 팔이 연직 위. 평면 모델의 두 관절축을 모두 A2 축 방향으로 잡으면 "
                             "q1 = A2, q2 = A4 * A4_axis_sign_relative_to_A2. 토크도 같은 부호 변환을 적용한다."),
        "link1": {"length": round(float((pos[3] - pos[1]) @ arm), 5), **seg1},
        "link2": {"length_to_tcp": round(float((tcp - pos[3]) @ arm), 5),
                  "length_to_flange": round(float((flange - pos[3]) @ arm), 5), **seg2},
        "payload": {"mass_spent_bed": params["payload"]["bed_spent_mass"],
                    "mass_new_bed": params["payload"]["bed_new_mass"],
                    "com_from_A4_along_link": round(float((bed_c - pos[3]) @ arm), 5),
                    "note": "운반 중(S3, S4 일부)에만 link2 끝에 더하는 점질량"},
        "torque_limit": {"A2": params["robot"]["torque_limit"][1], "A4": params["robot"]["torque_limit"][3]},
        "velocity_limit": {"A2": params["robot"]["velocity_limit"][1], "A4": params["robot"]["velocity_limit"][3]},
    }
